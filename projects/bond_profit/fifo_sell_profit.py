#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从两份原始 Excel 计算债券卖出结算差，FIFO，Decimal 精确到分。

python fifo_sell_profit.py  # 弹出文件选择窗口
python fifo_sell_profit.py 交易查询.xlsx 持仓记录.xlsx
python fifo_sell_profit.py --trades 交易查询.xlsx --holdings 持仓记录.xlsx --date 2026-09-24

默认纳入成交确认及交易执行中的买入和卖出；--confirmed-only 可切换口径。
买入按结算日/交易时间排序。
持仓数量视为本次卖出前可用数量；同日买入不重复导入持仓表已载明的批次。
部分卖出成本采用累计比例四舍五入之差，整批耗尽时成本与原结算金额一致。
输出默认为本脚本所在目录；按交易日期命名，相同日期再次运行会更新结果。
交易面额合计统计交易查询文件中所有交易（不按方向、日期或状态筛选），排除表尾合计行。
票息总额按用户指定的亿元口径读取原值，独立展示，不叠加到卖出结算差。
读取仅依赖标准库，写 Excel 需要 openpyxl，GUI 使用 Python 标准库 tkinter。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import posixpath
import re
import queue
import sys
import threading
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
RNS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
ZERO, CENT, WAN = Decimal('0'), Decimal('.01'), Decimal('10000')


def money(value):
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def decimal(value):
    result = Decimal(str(value).replace(',', '').strip())
    if not result.is_finite():
        raise ValueError(f'非有限数值: {value}')
    return result


def code(value):
    # 不转换成 int，保留债券代码前导零；仅去掉已知市场后缀。
    return re.sub(r'\.(SH|SZ|IB)$', '', str(value).strip(), flags=re.I)


def parse_date(value):
    text = str(value).strip()
    if re.fullmatch(r'\d+(\.\d+)?', text):
        return date(1899, 12, 30) + timedelta(days=int(decimal(text)))
    return date.fromisoformat(text[:10].replace('/', '-'))


def read_records(path, required):
    """忽略 WPS 非标准样式，只读 OOXML 的值和公式缓存，保留真实行号。"""
    results = []
    with zipfile.ZipFile(path) as z:
        shared = []
        if 'xl/sharedStrings.xml' in z.namelist():
            shared = [''.join(t.text or '' for t in si.iter(NS + 't'))
                      for si in ET.fromstring(z.read('xl/sharedStrings.xml'))]
        rels = {r.get('Id'): r.get('Target')
                for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        for sheet in ET.fromstring(z.read('xl/workbook.xml')).iter(NS + 'sheet'):
            target = rels[sheet.get(RNS + 'id')]
            xml_path = (target.lstrip('/') if target.startswith('/') else
                        posixpath.normpath(posixpath.join('xl', target)))
            headers = None
            for row in ET.fromstring(z.read(xml_path)).iter(NS + 'row'):
                cells = {}
                for c in row.findall(NS + 'c'):
                    v = c.find(NS + 'v')
                    value = '' if v is None else (v.text or '')
                    if c.get('t') == 's' and value:
                        value = shared[int(value)]
                    elif c.get('t') == 'inlineStr':
                        value = ''.join(t.text or '' for t in c.iter(NS + 't'))
                    col = re.sub(r'\d', '', c.get('r', ''))
                    cells[col] = value.strip()
                if headers is None:
                    if required.issubset(set(cells.values())):
                        headers = {c: v for c, v in cells.items() if v}
                    continue
                record = {headers[c]: v for c, v in cells.items() if c in headers}
                if any(record.values()):
                    record['_sheet'], record['_row'] = sheet.get('name'), int(row.get('r'))
                    results.append(record)
    if not results:
        raise ValueError(f'未找到字段 {sorted(required)}: {path}')
    return results


def location(path, row):
    return {'文件': str(path.resolve()), '工作表': row['_sheet'], '行号': row['_row']}


def infer_report_date(trades_path):
    records = read_records(trades_path, {'交易日期', '交易方向'})
    dates = {parse_date(r['交易日期']) for r in records
             if r.get('交易方向') == '卖出' and r.get('交易日期')}
    if len(dates) != 1:
        raise ValueError('无法识别唯一卖出交易日，请通过 --date 或 GUI 指定统计日期')
    return next(iter(dates))


def read_extra_metrics(trades_path, holdings_path):
    """所有业务的交易面额相加，包含文件内其他日期，不重复计入合计行。"""
    records = read_records(trades_path, {'交易编号', '交易日期', '交易方向', '交易面额(元)'})
    face_total, sources, ids = ZERO, [], set()
    for r in records:
        if not all(r.get(k) for k in ('交易编号', '交易日期', '交易方向')):
            continue
        if r['交易编号'] in ids:
            raise ValueError(f"交易面额合计发现重复交易编号: {r['交易编号']}")
        ids.add(r['交易编号'])
        amount = decimal(r['交易面额(元)'])
        if amount < ZERO:
            raise ValueError(f"交易面额为负: {r['交易编号']}")
        face_total += amount
        sources.append({'交易编号': r['交易编号'], '交易方向': r['交易方向'],
                        '交易日期': r['交易日期'], '交易面额元': amount,
                        '来源': location(trades_path, r)})
    if not sources:
        raise ValueError('交易查询没有可统计的交易记录')
    coupons = [r for r in read_records(holdings_path, {'票息总额'}) if r.get('票息总额')]
    if len(coupons) != 1:
        raise ValueError('票息总额必须有且仅有一个有效值，请检查源文件（公式需要保存计算缓存）')
    coupon = decimal(coupons[0]['票息总额'])
    return {'所有交易面额合计亿元': face_total / Decimal('100000000'), '票息总额亿元': coupon}, {
        '交易面额汇总': {'记录数': len(sources), '合计元': face_total,
                         '范围': '文件内全部交易，包含同业拆借及其他交易日；不含表尾合计行', '明细': sources},
        '票息总额': {'原值': coupon, '单位': '亿元', '单位依据': '用户指定，原值直接读取，不换算',
                     '来源': location(holdings_path, coupons[0])}}


def load_data(trades_path, holdings_path, target, confirmed_only=False, holding_only=False):
    trades_raw = read_records(trades_path, {'交易方向', '代码', '结算金额(元)'})
    holdings_raw = read_records(holdings_path, {'代码', '买入数量（万）', '结算金额（元）'})
    trades, lots, events = [], [], []
    seen = set()
    for r in trades_raw:
        direction = r.get('交易方向')
        if direction not in ('卖出', '买入', '分销买入') or not r.get('代码'):
            continue
        if parse_date(r['交易日期']) != target:
            continue
        status = r['交易状态']
        if status not in ('成交确认', '交易执行中'):
            events.append({'类型': '跳过交易状态', '交易状态': status, '来源': location(trades_path, r)})
            continue
        if confirmed_only and status != '成交确认':
            continue
        key = r['交易编号']
        if key in seen:
            raise ValueError(f'重复交易编号 {key}，不能自动重复计数')
        seen.add(key)
        qty, amount = decimal(r['交易面额(万)']), money(decimal(r['结算金额(元)']))
        if qty <= ZERO or amount < ZERO:
            raise ValueError(f'非法数量或结算金额: {key}')
        if decimal(r['交易面额(元)']) != qty * WAN:
            raise ValueError(f'万元/元面额不一致: {key}')
        when = datetime.fromisoformat(r['交易时间'])
        trades.append({'id': key, 'code': code(r['代码']), 'name': r['名称'],
                       'direction': direction, 'status': status, 'date': target,
                       'when': when, 'qty': qty, 'amount': amount,
                       'source': location(trades_path, r)})
    for r in holdings_raw:
        if not r.get('代码') or r.get('简称', '').startswith('【'):
            continue
        qty, amount = decimal(r['买入数量（万）']), money(decimal(r['结算金额（元）']))
        bought = parse_date(r['买入结算日'])
        if qty <= ZERO or amount < ZERO:
            raise ValueError(f'非法持仓数量/成本: {location(holdings_path, r)}')
        if bought > target:
            continue
        if bought == target:
            events.append({'类型': '持仓当日批次无成交时间', '说明': '按买入结算日及原始行号排序，视为当日开盘可用',
                           '来源': location(holdings_path, r)})
        lots.append({'id': f"持仓:{r['_sheet']}:{r['_row']}", 'code': code(r['代码']),
                     'name': r['简称'], 'qty': qty, 'amount': amount, 'remaining': qty,
                     'allocated': ZERO, 'date': bought, 'when': datetime.combine(bought, datetime.min.time()),
                     'type': '持仓记录', 'source': location(holdings_path, r)})
    same_day_signatures = defaultdict(int)
    for lot in lots:
        if lot['date'] == target:
            same_day_signatures[(lot['code'], lot['qty'], lot['amount'])] += 1
    if not holding_only:
        for t in trades:
            if t['direction'] == '卖出':
                continue
            signature = (t['code'], t['qty'], t['amount'])
            if same_day_signatures[signature] > 0:
                same_day_signatures[signature] -= 1
                events.append({'类型': '同日买入批次去重', '交易编号': t['id'], '来源': t['source']})
                continue
            lots.append({'id': t['id'], 'code': t['code'], 'name': t['name'], 'qty': t['qty'],
                         'amount': t['amount'], 'remaining': t['qty'], 'allocated': ZERO,
                         'date': target, 'when': t['when'], 'type': '当日买入', 'source': t['source']})
    return sorted(trades, key=lambda t: (t['when'], t['source']['行号'])), lots, events, holdings_raw


def fifo_match(sells, lots):
    """单条卖出可拆多批，重复卖出继续扣减余额，不允许未来买入补过去卖出。"""
    by_code = defaultdict(list)
    for lot in sorted(lots, key=lambda l: (l['when'], l['source']['行号'], l['id'])):
        by_code[lot['code']].append(lot)
    details, crossing, missing = [], [], []
    for sell in sorted(sells, key=lambda s: (s['when'], s['source']['行号'])):
        remaining, cost, matches = sell['qty'], ZERO, []
        eligible = [l for l in by_code[sell['code']] if l['remaining'] > ZERO and l['when'] <= sell['when']]
        first_remaining = eligible[0]['remaining'] if eligible else ZERO
        for lot in eligible:
            take = min(remaining, lot['remaining'])
            before = lot['remaining']
            used_after = lot['qty'] - before + take
            # 累计分摊，避免一批分成多笔时因逐笔四舍五入造成尾差。
            allocated_after = money(lot['amount'] * used_after / lot['qty'])
            part = allocated_after - lot['allocated']
            lot['allocated'], lot['remaining'] = allocated_after, before - take
            remaining, cost = remaining - take, cost + part
            matches.append({'买入编号': lot['id'], '买入来源': lot['type'], '买入结算日': lot['date'],
                            '买入记录数量万': lot['qty'], '买入结算金额元': lot['amount'],
                            '买入结算单价元每百元面额': lot['amount'] / (lot['qty'] * Decimal('100')),
                            '匹配前剩余数量万': before, '本次匹配数量万': take, '本次买入成本元': part,
                            '匹配后剩余数量万': lot['remaining'], '来源': lot['source']})
            if remaining == ZERO:
                break
        if remaining > ZERO:
            missing.append({'交易编号': sell['id'], '债券代码': sell['code'], '债券名称': sell['name'],
                            '未匹配数量万': remaining, '来源': sell['source'],
                            '说明': '截至卖出时间买入数量不足，完整盈利不计算，不将缺失成本当作零'})
        if eligible and sell['qty'] > first_remaining:
            crossing.append({'交易编号': sell['id'], '债券代码': sell['code'], '债券名称': sell['name'],
                             '卖出数量万': sell['qty'], '首条买入剩余数量万': first_remaining,
                             '涉及买入记录数': len(matches), '未匹配数量万': remaining,
                             '匹配明细': matches})
        details.append({'交易日期': sell['date'], '交易编号': sell['id'], '债券代码': sell['code'],
                        '债券名称': sell['name'], '交易状态': sell['status'], '卖出数量万': sell['qty'],
                        '卖出结算金额元': sell['amount'], '匹配买入成本元': cost,
                        '未匹配数量万': remaining,
                        '结算金额差元': sell['amount'] - cost if remaining == ZERO else None,
                        '来源': sell['source'], '匹配明细': matches})
    return details, crossing, missing


def aggregate(details):
    grouped = {}
    for d in details:
        b = grouped.setdefault(d['债券代码'], {'交易日期': d['交易日期'], '债券代码': d['债券代码'],
            '债券名称': d['债券名称'], '卖出笔数': 0, '卖出数量万': ZERO, '卖出结算金额元': ZERO,
            '匹配买入成本元': ZERO, '未匹配数量万': ZERO, '结算金额差元': ZERO})
        b['卖出笔数'] += 1
        for key in ('卖出数量万', '卖出结算金额元', '匹配买入成本元', '未匹配数量万'):
            b[key] += d[key]
        if d['结算金额差元'] is None or b['结算金额差元'] is None:
            b['结算金额差元'] = None
        else:
            b['结算金额差元'] += d['结算金额差元']
    return list(grouped.values())


def json_value(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise TypeError(type(value).__name__)


def write_excel(path, bonds, details, metrics, target, confirmed_only=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    summary, trace = wb.active, wb.create_sheet('FIFO匹配明细')
    summary.title = '当日盈利情况'
    summary.append(['交易日期', '债券代码', '债券名称', '结算金额差（元）', '说明', '汇总项目', '汇总值'])
    note = '卖出结算金额−FIFO按数量分摊的买入结算金额'
    for b in bonds:
        summary.append([target, b['债券代码'], b['债券名称'],
                        float(b['结算金额差元']) if b['结算金额差元'] is not None else '成本缺失',
                        note if summary.max_row == 1 else None])
    summary['E3'] = ('仅统计成交确认，交易执行中全部排除' if confirmed_only else
                     '统计成交确认及交易执行中的买入和卖出')
    summary['E4'] = '持仓表无记录时，采用卖出前的当日买入成本'
    total_row = len(bonds) + 2
    summary.cell(total_row, 1, '合计')
    total = metrics['当日结算金额差合计元']
    summary.cell(total_row, 4, float(total) if total is not None else '成本缺失，无法合计')
    for i, (label, value) in enumerate(metrics.items(), 2):
        summary.cell(i, 6, label)
        summary.cell(i, 7, float(value) if isinstance(value, Decimal) else
                     ('成本缺失，无法合计' if value is None else value))
    trace.append(['卖出交易编号', '卖出状态', '债券代码', '债券名称', '卖出数量（万）',
                  '卖出结算金额（元）', '匹配买入成本（元）', '结算金额差（元）',
                  '买入来源', '买入编号', '买入结算日', '买入原数量（万）', '买入结算金额（元）',
                  '本次匹配数量（万）', '本次成本（元）', '批次剩余数量（万）', '买入来源行号', '卖出来源行号'])
    for d in details:
        for i, m in enumerate(d['匹配明细'] or [None]):
            row = [d['交易编号'], d['交易状态'], d['债券代码'], d['债券名称']]
            row += ([d['卖出数量万'], d['卖出结算金额元'], d['匹配买入成本元'],
                     d['结算金额差元'] if d['结算金额差元'] is not None else '成本缺失']
                    if i == 0 else [None] * 4)
            row += ([m['买入来源'], m['买入编号'], m['买入结算日'], m['买入记录数量万'],
                     m['买入结算金额元'], m['本次匹配数量万'], m['本次买入成本元'],
                     m['匹配后剩余数量万'], m['来源']['行号']] if m else [None] * 9)
            row += [d['来源']['行号']]
            trace.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = 'A2'
        for row in ws:
            for cell in row:
                cell.font = Font(name='Microsoft YaHei', size=11)
                cell.alignment = Alignment(vertical='center', horizontal='right' if isinstance(cell.value, (int, float)) else 'left')
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
                if isinstance(cell.value, date):
                    cell.number_format = 'yyyy-mm-dd'
            ws.row_dimensions[row[0].row].height = 27
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='203864')
            cell.font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 44
    for i, width in enumerate((15, 15, 33, 24, 53, 36, 25), 1):
        summary.column_dimensions[get_column_letter(i)].width = width
    for i, (label, value) in enumerate(metrics.items(), 2):
        summary.cell(i, 7).number_format = ('0' if isinstance(value, int) else
                                          ('0.00########' if label.endswith('亿元') else '#,##0.00'))
    for i, width in enumerate((18, 17, 16, 33, 19, 24, 24, 23, 18, 23, 17, 22, 26, 22, 24, 24, 18, 18), 1):
        trace.column_dimensions[get_column_letter(i)].width = width
    for row in trace.iter_rows(min_row=2, min_col=17, max_col=18):
        for cell in row:
            cell.number_format = '0'
    for cell in summary[total_row][:4]:
        cell.font = Font(name='Microsoft YaHei', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='E2EAF4')
    summary.print_options.horizontalCentered = True
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = 'landscape'
    summary.page_setup.paperSize = summary.PAPERSIZE_A3
    summary.page_setup.fitToWidth, summary.page_setup.fitToHeight = 1, 1
    summary.print_area = f'A1:G{max(total_row, len(metrics)+1)}'
    wb.save(path)


def run_analysis(trades_path, holdings_path, report_date=None, output_dir=None,
                 confirmed_only=False, holding_only=False):
    trades_path, holdings_path = Path(trades_path).resolve(), Path(holdings_path).resolve()
    output_dir = Path(output_dir).resolve() if output_dir else ROOT
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        target = parse_date(report_date) if report_date else infer_report_date(trades_path)
        trades, lots, events, _ = load_data(trades_path, holdings_path, target, confirmed_only, holding_only)
        extra_metrics, extra_sources = read_extra_metrics(trades_path, holdings_path)
        sells = [t for t in trades if t['direction'] == '卖出']
        details, crossing, missing = fifo_match(sells, lots)
        bonds = aggregate(details)
        total_profit = sum((d['结算金额差元'] for d in details), ZERO) if not missing else None
        metrics = {'当日结算金额差合计元': total_profit,
                   '卖出结算金额合计元': sum((t['amount'] for t in sells), ZERO),
                   '匹配买入成本合计元': sum((d['匹配买入成本元'] for d in details), ZERO),
                   '当日债券买卖结算金额合计亿元': sum((t['amount'] for t in trades), ZERO) / Decimal('100000000'),
                   '卖出笔数': len(sells), '卖出债券只数': len(bonds),
                   '超过首条买入剩余数量事件数': len(crossing), '未匹配卖出笔数': len(missing)}
        metrics.update(extra_metrics)
        for status in ('成交确认', '交易执行中'):
            subset = [d for d in details if d['交易状态'] == status]
            metrics[f'{status}卖出结算差元'] = (
                sum((d['结算金额差元'] for d in subset), ZERO)
                if all(d['结算金额差元'] is not None for d in subset) else None)
        used_by_lot = defaultdict(list)
        for d in details:
            for m in d['匹配明细']:
                used_by_lot[m['买入编号']].append(d['交易编号'])
        log = {'规则': {'交易日': target, '买卖状态范围': ('仅成交确认，交易执行中全部排除' if confirmed_only else
                                                      '成交确认及交易执行中的买入和卖出'),
                       '买入来源': '仅持仓表' if holding_only else '持仓表及交易查询当日买入',
                       'FIFO排序': '买入结算日及原行号；当日买卖按交易时间及原行号',
                       '分摊规则': '按匹配数量分摊结算金额，累计成本四舍五入到分；盈利为结算差，不另加票息或费用',
                       '汇报条件': '单笔卖出数量超过FIFO首条可用买入的剩余数量；数量不足不计算完整盈利',
                       '持仓假设': '表内买入数量视为卖出前可用数量，代投代缴批次也参与匹配'},
               '输入文件': [{'路径': str(p.resolve()), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in (trades_path, holdings_path)],
               '新增指标来源': extra_sources,
               '汇总': metrics, '卖出跨多条买入记录': crossing, '未匹配卖出': missing,
               '读取及去重日志': events,
               '持仓表缺失而使用当日买入': [d for d in details if any(m['买入来源'] == '当日买入' for m in d['匹配明细'])],
               '同一买入批次匹配多笔卖出': [{'买入编号': k, '卖出笔数': len(v), '卖出交易编号': v} for k, v in used_by_lot.items() if len(v) > 1],
               '债券汇总': bonds, '卖出明细': details,
               '已使用批次剩余数量': [{'买入编号': l['id'], '债券代码': l['code'], '剩余数量万': l['remaining'],
                                        '已分摊成本元': l['allocated'], '来源': l['source']} for l in lots if l['allocated'] != ZERO]}
        # 数据守恒：匹配数量 + 缺口 == 卖出量；已分摊成本不得超过批次总成本。
        for d in details:
            assert sum((m['本次匹配数量万'] for m in d['匹配明细']), ZERO) + d['未匹配数量万'] == d['卖出数量万']
            assert sum((m['本次买入成本元'] for m in d['匹配明细']), ZERO) == d['匹配买入成本元']
        for l in lots:
            assert ZERO <= l['remaining'] <= l['qty'] and ZERO <= l['allocated'] <= l['amount']
            if l['remaining'] == ZERO:
                assert l['allocated'] == l['amount']
        if total_profit is not None:
            assert total_profit == metrics['卖出结算金额合计元'] - metrics['匹配买入成本合计元']
            assert total_profit == sum((b['结算金额差元'] for b in bonds), ZERO)
        log_path = output_dir / f'{target}_FIFO日志.json'
        xlsx_path = output_dir / f'{target}_FIFO盈利统计.xlsx'
        csv_path = output_dir / f'{target}_FIFO盈利汇总.csv'
        log_path.write_text(json.dumps(log, ensure_ascii=False, indent=2, default=json_value), encoding='utf-8')
        with csv_path.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['交易日期', '债券代码', '债券名称', '结算金额差（元）'])
            writer.writerows([[target, b['债券代码'], b['债券名称'], b['结算金额差元']] for b in bonds])
            writer.writerow(['合计', '', '', total_profit])
            writer.writerow([])
            writer.writerow(['汇总项目', '汇总值'])
            writer.writerows(metrics.items())
        write_excel(xlsx_path, bonds, details, metrics, target, confirmed_only)
        return {'汇总': metrics, '债券汇总': bonds, '输出目录': str(output_dir),
                'Excel': str(xlsx_path), 'JSON': str(log_path), 'CSV': str(csv_path),
                '退出码': 2 if missing else 0}
    except Exception as exc:
        (output_dir / '运行错误.json').write_text(json.dumps({'错误类型': type(exc).__name__, '错误': str(exc)}, ensure_ascii=False, indent=2), encoding='utf-8')
        raise


def build_gui(initial=None):
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    initial = initial or {}
    root = tk.Tk()
    root.title('债券 FIFO 盈利统计')
    root.geometry('650x300')
    root.resizable(False, False)
    panel = ttk.Frame(root, padding=16)
    panel.pack(fill='both', expand=True)
    panel.columnconfigure(1, weight=1)
    trades_var = tk.StringVar(value=str(initial.get('trades') or ''))
    holdings_var = tk.StringVar(value=str(initial.get('holdings') or ''))
    date_var = tk.StringVar(value=initial.get('date') or '')
    confirmed_var = tk.BooleanVar(value=initial.get('confirmed_only', False))
    status_var = tk.StringVar(value='请选择两份 Excel 文件。')
    messages = queue.SimpleQueue()
    running = False
    def choose(var):
        path = filedialog.askopenfilename(parent=root, title='选择 Excel 文件',
                                         filetypes=[('Excel 工作簿', '*.xlsx')])
        if path:
            var.set(path)
    for i, (label, var) in enumerate((('交易查询文件', trades_var), ('持仓记录文件', holdings_var))):
        ttk.Label(panel, text=label).grid(row=i, column=0, padx=(0, 12), pady=8, sticky='w')
        ttk.Entry(panel, textvariable=var, width=50).grid(row=i, column=1, sticky='ew')
        ttk.Button(panel, text='选择文件', command=lambda v=var: choose(v)).grid(row=i, column=2, padx=(8, 0))
    ttk.Label(panel, text='统计日期').grid(row=2, column=0, pady=8, sticky='w')
    ttk.Entry(panel, textvariable=date_var).grid(row=2, column=1, sticky='ew')
    ttk.Label(panel, text='留空自动识别').grid(row=2, column=2, padx=(8, 0))
    ttk.Checkbutton(panel, text='仅成交确认（默认也统计交易执行中）', variable=confirmed_var).grid(row=3, column=1, sticky='w', pady=8)
    ttk.Label(panel, text=f'输出目录：{initial.get("output_dir") or ROOT}', wraplength=610).grid(row=4, column=0, columnspan=3, sticky='w')
    ttk.Label(panel, textvariable=status_var, wraplength=610).grid(row=6, column=0, columnspan=3, pady=8, sticky='w')
    def start():
        nonlocal running
        if running:
            return
        if not trades_var.get().strip() or not holdings_var.get().strip():
            messagebox.showerror('文件未选择', '请选择交易查询和持仓记录文件。', parent=root)
            return
        params = (trades_var.get().strip(), holdings_var.get().strip(), date_var.get().strip() or None,
                  initial.get('output_dir'), confirmed_var.get(), initial.get('holding_only', False))
        running = True
        start_button.configure(state='disabled')
        status_var.set('正在计算并生成文件…')
        def worker():
            try:
                messages.put(('ok', run_analysis(*params)))
            except Exception as exc:
                messages.put(('error', str(exc)))
        threading.Thread(target=worker, daemon=True).start()
    def poll():
        nonlocal running
        try:
            kind, result = messages.get_nowait()
        except queue.Empty:
            root.after(100, poll)
            return
        running = False
        start_button.configure(state='normal')
        if kind == 'error':
            status_var.set('计算失败，请检查文件和日期。')
            messagebox.showerror('计算失败', result, parent=root)
        else:
            profit = result['汇总']['当日结算金额差合计元']
            status_var.set(f'完成：盈利 {profit if profit is not None else "成本缺失"} 元，文件已保存。')
            messagebox.showinfo('已生成', f"Excel 和 JSON 已保存到：\n{result['输出目录']}", parent=root)
        root.after(100, poll)
    start_button = ttk.Button(panel, text='生成汇总和日志', command=start)
    start_button.grid(row=5, column=1, pady=10, sticky='w')
    def close():
        if running:
            status_var.set('正在写入结果，请完成后关闭。')
        else:
            root.destroy()
    root.protocol('WM_DELETE_WINDOW', close)
    root.after(100, poll)
    return root


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('trade_file', type=Path, nargs='?', help='交易查询文件（文件名可任意）')
    parser.add_argument('holding_file', type=Path, nargs='?', help='持仓记录文件（文件名可任意）')
    parser.add_argument('--trades', type=Path)
    parser.add_argument('--holdings', type=Path)
    parser.add_argument('--date', help='YYYY-MM-DD，留空自动识别唯一卖出交易日')
    parser.add_argument('--output-dir', type=Path, help='默认与 Python 脚本同目录')
    parser.add_argument('--confirmed-only', action='store_true')
    parser.add_argument('--holding-only', action='store_true')
    parser.add_argument('--gui', action='store_true', help='弹出文件选择窗口')
    args = parser.parse_args(argv)
    if (args.trade_file and args.trades) or (args.holding_file and args.holdings):
        parser.error('请使用位置参数或 --trades/--holdings，不要重复指定文件')
    trades, holdings = args.trades or args.trade_file, args.holdings or args.holding_file
    if args.gui or (trades is None and holdings is None):
        try:
            app = build_gui({'trades': trades, 'holdings': holdings, 'date': args.date,
                             'output_dir': args.output_dir, 'confirmed_only': args.confirmed_only,
                             'holding_only': args.holding_only})
            app.mainloop()
            return 0
        except Exception as exc:
            print(f'无法启动 GUI: {exc}。请安装含 tkinter 的 Python，或在终端提供两份文件。', file=sys.stderr)
            return 1
    if trades is None or holdings is None:
        parser.error('终端模式需要同时提供交易查询和持仓记录文件')
    try:
        result = run_analysis(trades, holdings, args.date, args.output_dir,
                              args.confirmed_only, args.holding_only)
        print(json.dumps(result, ensure_ascii=False, indent=2, default=json_value))
        return result['退出码']
    except Exception as exc:
        print(f'计算失败：{exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())

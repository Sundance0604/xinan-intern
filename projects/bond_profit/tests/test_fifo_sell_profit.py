"""python -m unittest discover -s projects/bond_profit/tests -v"""
import unittest
import contextlib
import io
import json
import shutil
import sys
import tempfile
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fifo_sell_profit import ROOT, ZERO, aggregate, fifo_match, load_data, read_extra_metrics, run_analysis, main

LOCAL_TRADES = ROOT/'local_data/0924交易查询.xlsx'
LOCAL_HOLDINGS = ROOT/'local_data/债券持仓及舆情分析_20260924持仓记录.xlsx'


def lot(key, qty, amount, when='2026-09-01 00:00:00', row=2):
    return {'id': key, 'code': '032000001', 'name': '测试债券', 'qty': D(qty),
            'amount': D(amount), 'remaining': D(qty), 'allocated': ZERO,
            'when': datetime.fromisoformat(when), 'date': date.fromisoformat(when[:10]),
            'type': '持仓记录', 'source': {'行号': row}}


def sell(key, qty, amount, when='2026-09-24 12:00:00', row=2):
    return {'id': key, 'code': '032000001', 'name': '测试债券', 'qty': D(qty),
            'amount': D(amount), 'when': datetime.fromisoformat(when),
            'date': date.fromisoformat(when[:10]), 'status': '成交确认', 'source': {'行号': row}}


class FifoTests(unittest.TestCase):
    @unittest.skipUnless(LOCAL_TRADES.exists(), '本地真实数据不随仓库发布')
    def test_default_includes_executing_buys_and_sells(self):
        trades, lots, _, _ = load_data(LOCAL_TRADES, LOCAL_HOLDINGS, date(2026, 9, 24))
        ds, cross, missing = fifo_match([t for t in trades if t['direction'] == '卖出'], lots)
        self.assertEqual((len(ds), len(aggregate(ds)), len(cross), len(missing)), (12, 9, 0, 0))
        self.assertEqual(sum((d['结算金额差元'] for d in ds), ZERO), D('1931456.38'))
        self.assertEqual(sum((d['结算金额差元'] for d in ds if d['交易状态']=='交易执行中'), ZERO), D('289193.77'))
        self.assertEqual(sum((t['amount'] for t in trades), ZERO), D('1124571959.10'))
        self.assertTrue(any(l['id'] == '443431' for l in lots))
        self.assertTrue(any(l['id'] == '443600' for l in lots))

    def test_cross_lots_after_previous_partial_sale(self):
        # 输入顺序故意反向：先扣最早批次，第二次卖出跨越余下的 40 和下批的 40。
        lots = [lot('B', '50', '600', '2026-09-02 00:00:00'), lot('A', '100', '1000')]
        ds, crossed, missing = fifo_match([sell('S2', '80', '1000', row=3), sell('S1', '60', '700')], lots)
        self.assertEqual([d['交易编号'] for d in ds], ['S1', 'S2'])
        self.assertEqual([d['匹配买入成本元'] for d in ds], [D('600'), D('880')])
        self.assertEqual([d['结算金额差元'] for d in ds], [D('100'), D('120')])
        self.assertEqual(crossed[0]['首条买入剩余数量万'], D('40'))
        self.assertEqual(crossed[0]['涉及买入记录数'], 2)
        self.assertEqual(missing, [])
        self.assertEqual(lots[0]['remaining'], D('10'))

    def test_rounding_preserves_full_lot_cost(self):
        ds, crossed, missing = fifo_match([sell(str(i), '1', '1', row=i+2) for i in range(3)], [lot('A', '3', '1.00')])
        self.assertEqual([d['匹配买入成本元'] for d in ds], [D('.33'), D('.34'), D('.33')])
        self.assertEqual(sum((d['匹配买入成本元'] for d in ds), ZERO), D('1.00'))
        self.assertEqual(crossed, [])
        self.assertEqual(missing, [])

    def test_future_buy_and_missing_cost_do_not_create_profit(self):
        ds, crossed, missing = fifo_match([sell('S', '80', '1000')],
            [lot('old', '50', '500'), lot('future', '100', '900', '2026-09-24 13:00:00')])
        self.assertEqual(ds[0]['未匹配数量万'], D('30'))
        self.assertIsNone(ds[0]['结算金额差元'])
        self.assertIsNone(aggregate(ds)[0]['结算金额差元'])
        self.assertEqual(len(missing), 1)
        self.assertEqual(len(crossed), 1)
        empty, _, missing = fifo_match([sell('S', '80', '1000')], [])
        self.assertIsNone(empty[0]['结算金额差元'])
        self.assertEqual(missing[0]['未匹配数量万'], D('80'))

    @unittest.skipUnless(LOCAL_TRADES.exists(), '本地真实数据不随仓库发布')
    def test_source_workbooks_reconcile(self):
        trades, lots, _, _ = load_data(LOCAL_TRADES, LOCAL_HOLDINGS, date(2026, 9, 24), confirmed_only=True)
        sells = [t for t in trades if t['direction'] == '卖出']
        ds, cross, missing = fifo_match(sells, deepcopy(lots))
        self.assertEqual((len(ds), len(aggregate(ds)), len(cross), len(missing)), (8, 6, 0, 0))
        self.assertEqual(sum((d['结算金额差元'] for d in ds), ZERO), D('1642262.61'))
        self.assertEqual(sum((d['结算金额差元'] for d in ds if d['交易状态']=='成交确认'), ZERO), D('1642262.61'))
        # 只用持仓表时 2 只、2 笔卖出没有成本，不得把整笔卖款当盈利。
        strict, _, missing = fifo_match(sells, [l for l in deepcopy(lots) if l['type']=='持仓记录'])
        self.assertEqual(len(missing), 2)
        self.assertEqual(sum(d['结算金额差元'] is None for d in strict), 2)
        # 买卖结算额合计：排除同业拆借及尾部总计行。
        self.assertEqual(sum((t['amount'] for t in trades), ZERO), D('898853404.85'))
        self.assertTrue(all(t['status'] == '成交确认' for t in trades))
        self.assertTrue(all(l['id'] not in ('443431', '443600') for l in lots))

    @unittest.skipUnless(LOCAL_TRADES.exists(), '本地真实数据不随仓库发布')
    def test_all_face_volume_and_coupon_original_value(self):
        metrics, sources = read_extra_metrics(LOCAL_TRADES, LOCAL_HOLDINGS)
        self.assertEqual(metrics['所有交易面额合计亿元'], D('21.25'))
        self.assertEqual(metrics['票息总额亿元'], D('17.11123288'))
        self.assertEqual(sources['交易面额汇总']['记录数'], 31)
        self.assertEqual(sources['票息总额']['来源']['行号'], 2)

    def test_arbitrary_filenames_date_inference_and_cli(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            trades, holdings = base/'任何交易文件.xlsx', base/'不同持仓名称.xlsx'
            shutil.copyfile(ROOT/'test_data/sample_trades.xlsx', trades)
            shutil.copyfile(ROOT/'test_data/sample_holdings.xlsx', holdings)
            result = run_analysis(trades, holdings, output_dir=base)
            self.assertEqual(result['汇总']['当日结算金额差合计元'], D('55000.00'))
            self.assertEqual(result['汇总']['所有交易面额合计亿元'], D('.041'))
            self.assertEqual(result['汇总']['票息总额亿元'], D('.0000025'))
            self.assertEqual(result['汇总']['超过首条买入剩余数量事件数'], 1)
            self.assertEqual(result['退出码'], 0)
            self.assertEqual(Path(result['Excel']).parent, base)
            self.assertEqual(Path(result['JSON']).parent, base)
            log = json.loads(Path(result['JSON']).read_text(encoding='utf-8'))
            self.assertEqual(log['规则']['交易日'], '2026-09-24')
            ws = load_workbook(result['Excel'], data_only=True)['当日盈利情况']
            values = {ws.cell(i, 6).value: ws.cell(i, 7).value for i in range(2, ws.max_row+1)}
            self.assertEqual(values['所有交易面额合计亿元'], .041)
            self.assertEqual(values['票息总额亿元'], .0000025)
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main([str(trades), str(holdings), '--output-dir', str(base)]), 0)

    def test_no_arguments_opens_gui_and_default_output_is_script_directory(self):
        from unittest.mock import Mock, patch
        import fifo_sell_profit
        gui = Mock()
        with patch.object(fifo_sell_profit, 'build_gui', return_value=gui) as builder:
            self.assertEqual(main([]), 0)
            builder.assert_called_once()
            gui.mainloop.assert_called_once()
        self.assertEqual(ROOT, Path(fifo_sell_profit.__file__).resolve().parent)

    def test_terminal_requires_both_files(self):
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as error:
            main(['--trades', 'missing.xlsx'])
        self.assertEqual(error.exception.code, 2)

    def test_missing_coupon_is_reported_in_error_json(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            holdings = base/'blank_coupon.xlsx'
            wb = load_workbook(ROOT/'test_data/sample_holdings.xlsx')
            wb.active['F2'] = None
            wb.save(holdings)
            with self.assertRaisesRegex(ValueError, '票息总额'):
                run_analysis(ROOT/'test_data/sample_trades.xlsx', holdings, output_dir=base)
            error = json.loads((base/'运行错误.json').read_text(encoding='utf-8'))
            self.assertIn('票息总额', error['错误'])
            self.assertFalse((base/'2026-09-24_FIFO盈利统计.xlsx').exists())

    def test_multiple_sell_dates_require_explicit_date(self):
        from openpyxl import load_workbook
        from fifo_sell_profit import infer_report_date
        with tempfile.TemporaryDirectory() as folder:
            trades = Path(folder)/'multi_dates.xlsx'
            wb = load_workbook(ROOT/'test_data/sample_trades.xlsx')
            wb.active['C3'] = '2026-09-25'
            wb.save(trades)
            with self.assertRaisesRegex(ValueError, '唯一卖出交易日'):
                infer_report_date(trades)


if __name__ == '__main__':
    unittest.main()

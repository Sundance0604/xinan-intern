#!/usr/bin/env python3
"""对比早盘/晚间持仓表，输出各区块新增、删除的行（只看行的增删，不比对数值）。

用法: python diff_holdings.py 早盘.xlsx 晚间.xlsx [sheet名，默认第一个sheet]
输出: JSON
"""
import sys, json
from collections import Counter
import openpyxl

MARK_HOLD = "【持仓】"
MARK_AGENT = "【代投代缴】"
# 用于识别"同一行"的字段（表头名）。这些字段买入后不会变化，其余列（估值、资本利得、剩余年限、持仓工作日等）每天会变。
KEY_COLS = ["代码", "买入收益率", "买入数量（万）", "买入结算日"]


def norm(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 6)
    return str(v).strip()


def load(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header = {str(c.value).strip().split("\n")[0]: i for i, c in enumerate(ws[1]) if c.value}
    # 表头名匹配（容忍括号/换行差异）
    def col(name):
        for h, i in header.items():
            if h.replace("(", "（").replace(")", "）") == name:
                return i
        raise KeyError(f"表头里找不到列: {name}，现有表头: {list(header)}")
    idx = {n: col(n) for n in KEY_COLS + ["简称", "担保/非担保", "结算金额（元）", "资本利得"]}
    section = None
    rows = {"持仓": [], "代投代缴": []}
    for r in ws.iter_rows(min_row=2):
        a = r[0].value
        a = str(a).strip() if a is not None else ""
        if a == MARK_HOLD:
            section = "持仓"; continue          # 第2行是汇总行，不是持仓
        if a == MARK_AGENT:
            section = "代投代缴"; continue
        if section is None:
            continue
        code = r[idx["代码"]].value
        if not code:                              # 代码为空 => 区块结束（敞口/拆借等汇总区）
            if section == "代投代缴":
                section = None
            continue
        key = tuple(norm(r[idx[c]].value) for c in KEY_COLS)
        rows[section].append({
            "row": r[0].row, "key": key, "简称": a,
            "担保": norm(r[idx["担保/非担保"]].value),
            "结算金额为空": norm(r[idx["结算金额（元）"]].value) == "",
            "资本利得为空": norm(r[idx["资本利得"]].value) in ("", "-"),
        })
    return rows


def diff(old, new):
    out = {}
    for sec in ("持仓", "代投代缴"):
        co = Counter(x["key"] for x in old[sec]); cn = Counter(x["key"] for x in new[sec])
        removed_k = co - cn; added_k = cn - co
        def pick(rows, ks):
            ks = Counter(ks); res = []
            for x in rows:
                if ks[x["key"]] > 0:
                    ks[x["key"]] -= 1; res.append(x)
            return res
        added = pick(new[sec], added_k); removed = pick(old[sec], removed_k)
        out[sec] = {
            "早盘行数": len(old[sec]), "晚间行数": len(new[sec]),
            "新增": [{"简称": x["简称"], "代码": x["key"][0], "买入收益率": x["key"][1], "买入数量(万)": x["key"][2],
                     "买入结算日": x["key"][3], "晚间表行号": x["row"], "担保": x["担保"],
                     "结算金额为空": x["结算金额为空"], "资本利得为空": x["资本利得为空"]} for x in added],
            "删除": [{"简称": x["简称"], "代码": x["key"][0], "买入收益率": x["key"][1], "买入数量(万)": x["key"][2],
                     "买入结算日": x["key"][3], "早盘表行号": x["row"], "担保": x["担保"]} for x in removed],
        }
    return out


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    sheet = sys.argv[3] if len(sys.argv) > 3 else None
    print(json.dumps(diff(load(sys.argv[1], sheet), load(sys.argv[2], sheet)), ensure_ascii=False, indent=1))

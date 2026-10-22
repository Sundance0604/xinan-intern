#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""对比当日与最近一个交易日的持仓表，输出「当日买入卖出记录」txt。

用法:
  python daily_trade_records.py [--date YYYY-MM-DD] [--today 当日持仓表.xlsx]
                                [--prev 上一交易日持仓表.xlsx] [--out 输出.txt]

默认:
  - 当日表 = 本地数据目录/持仓表-YYYY-MM-DD.xlsx（--date 默认今天）
  - 上一份 = 日期 < 当日的最近一份（周一自动落到上周五）
  - 输出   = 本地数据目录/当日买入卖出记录-YYYY-MM-DD.txt

差异口径:
  - 以「代码 + 买入收益率 + 买入数量（万）+ 买入结算日」标识一行
  - 当日表新增的行 = 当日买入；当日表消失的行 = 当日卖出
  - 类型：代投代缴区 → 一级买入/卖出；担保 → 担保买入/卖出；其余 → 二级买入/卖出
  - 金额：取该行「结算金额（元）」
"""
import sys, os, re, glob, argparse, datetime
from collections import Counter
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DESK = os.environ.get('SWSC_DATA_DIR', os.path.join(os.path.dirname(os.path.dirname(__file__)), 'local_data'))
KEY_COLS = ["代码", "买入收益率", "买入数量（万）", "买入结算日"]
COLS = ["简称", "代码", "买入收益率", "买入数量（万）", "结算金额（元）",
        "买入结算日", "主办人", "担保/非担保"]


def parse_date_file(base):
    m = re.search(r"持仓表[-_](\d{4})[-_]?(\d{2})[-_]?(\d{2})", base)
    if not m:
        return None
    return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def norm(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 6)
    return str(v).strip()


def num(v):
    if v is None:
        return 0.0
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "--"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def find_holdings(target):
    exact = os.path.join(DESK, f"持仓表-{target.strftime('%Y-%m-%d')}.xlsx")
    if os.path.exists(exact):
        return exact
    for f in sorted(glob.glob(os.path.join(DESK, "持仓表[-_]*.xlsx"))):
        if parse_date_file(os.path.basename(f)) == target:
            return f
    return None


def find_prev(target):
    best = None
    for f in glob.glob(os.path.join(DESK, "持仓表[-_]*.xlsx")):
        d = parse_date_file(os.path.basename(f))
        if d and d < target and (best is None or d > best[0]):
            best = (d, f)
    return best[1] if best else None


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header = {str(c.value).strip().split("\n")[0]: c.column for c in ws[1] if c.value}

    def col(name):
        for h, i in header.items():
            if h == name or h.replace("(", "（").replace(")", "）") == name:
                return i
        return None

    idx = {k: col(k) for k in COLS}
    missing = [k for k, v in idx.items() if v is None]
    if missing:
        raise KeyError(f"持仓表缺少列: {missing}，现有表头: {list(header)}")

    def cell(r, k):
        c = idx[k]
        return r[c - 1] if c else None

    rows = {"持仓": [], "代投代缴": []}
    section = None
    for r in ws.iter_rows(min_row=2):
        a = str(r[0].value).strip() if r[0].value is not None else ""
        if a == "【持仓】":
            section = "持仓"
            continue
        if a == "【代投代缴】":
            section = "代投代缴"
            continue
        if section is None:
            continue
        code = str(cell(r, "代码").value or "").strip()
        if not code:
            if section == "代投代缴":
                section = None
            continue
        rows[section].append({
            "区块": section,
            "简称": str(cell(r, "简称").value or "").strip(),
            "代码": code,
            "买入收益率": norm(cell(r, "买入收益率").value),
            "数量万": num(cell(r, "买入数量（万）").value),
            "结算金额": num(cell(r, "结算金额（元）").value),
            "买入结算日": norm(cell(r, "买入结算日").value),
            "主办人": str(cell(r, "主办人").value or "").strip(),
            "担保": str(cell(r, "担保/非担保").value or "").strip(),
            "key": tuple(norm(cell(r, k).value) for k in KEY_COLS),
        })
    return rows


def kind_of(row, direction):
    if row["区块"] == "代投代缴":
        return f"一级{direction}"
    if row["担保"].strip() == "担保":
        return f"担保{direction}"
    return f"二级{direction}"


def diff(old, new):
    added, removed = [], []
    for sec in ("持仓", "代投代缴"):
        co = Counter(x["key"] for x in old[sec])
        cn = Counter(x["key"] for x in new[sec])
        add_k, rem_k = cn - co, co - cn

        def pick(rows, ks):
            ks = Counter(ks)
            out = []
            for x in rows:
                if ks[x["key"]] > 0:
                    ks[x["key"]] -= 1
                    out.append(x)
            return out

        added.extend(pick(new[sec], add_k))
        removed.extend(pick(old[sec], rem_k))
    return added, removed


def money(x):
    return f"{x:,.2f}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default=None)
    ap.add_argument("--today", default=None)
    ap.add_argument("--prev", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    target = datetime.date.today()
    if args.date:
        target = datetime.datetime.strptime(args.date, "%Y-%m-%d").date()

    today_file = args.today or find_holdings(target)
    prev_file = args.prev or find_prev(target)
    if not today_file:
        sys.exit(f"找不到当日的持仓表：持仓表-{target.strftime('%Y-%m-%d')}.xlsx")
    if not prev_file:
        sys.exit("找不到上一交易日的持仓表（日期早于当日的最近一份）")

    new = load(today_file)
    old = load(prev_file)
    added, removed = diff(old, new)

    buys = [{**x, "类型": kind_of(x, "买入")} for x in added]
    sells = [{**x, "类型": kind_of(x, "卖出")} for x in removed]

    def order(x):
        return ({"一级买入": 0, "担保买入": 1, "二级买入": 2,
                 "一级卖出": 0, "担保卖出": 1, "二级卖出": 2}.get(x["类型"], 9),
                x["主办人"], x["代码"])

    buys.sort(key=order)
    sells.sort(key=order)

    lines = [
        f"{target.strftime('%Y-%m-%d')} 当日买入卖出记录",
        f"对比基准：{os.path.basename(prev_file)}  →  {os.path.basename(today_file)}",
        "",
        f"一、买入（{len(buys)} 笔，合计 {money(sum(x['结算金额'] for x in buys))} 元）",
    ]
    if not buys:
        lines.append("  （无）")
    for i, x in enumerate(buys, 1):
        lines.append(
            f"{i}. {x['类型']} | {x['简称']} | {x['代码']} | 主办人：{x['主办人'] or '未填'} | "
            f"数量：{x['数量万']:g}万 | 买入金额：{money(x['结算金额'])} 元 | 买入结算日：{x['买入结算日']}"
        )
    lines += [
        "",
        f"二、卖出（{len(sells)} 笔，合计 {money(sum(x['结算金额'] for x in sells))} 元）",
    ]
    if not sells:
        lines.append("  （无）")
    for i, x in enumerate(sells, 1):
        lines.append(
            f"{i}. {x['类型']} | {x['简称']} | {x['代码']} | 主办人：{x['主办人'] or '未填'} | "
            f"数量：{x['数量万']:g}万 | 卖出金额：{money(x['结算金额'])} 元 | 买入结算日：{x['买入结算日']}"
        )
    lines += [
        "",
        f"汇总：买入 {len(buys)} 笔 / {money(sum(x['结算金额'] for x in buys))} 元；"
        f"卖出 {len(sells)} 笔 / {money(sum(x['结算金额'] for x in sells))} 元",
        f"—— {target.strftime('%Y-%m-%d')}",
    ]

    text = "\n".join(lines) + "\n"
    out = args.out or os.path.join(DESK, f"当日买入卖出记录-{target.strftime('%Y-%m-%d')}.txt")
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)
    print(text)
    print(f"[已写入] {out}")


if __name__ == "__main__":
    main()

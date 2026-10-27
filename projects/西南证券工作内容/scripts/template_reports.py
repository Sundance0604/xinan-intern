"""Fill the user-approved report templates without replacing their formatting."""
from copy import copy
from decimal import Decimal
from pathlib import Path
import datetime
import json

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

TEMPLATES = Path(__file__).resolve().parents[1] / 'templates'


def clone_sheet(source, target):
    """Copy styles by value; style IDs are not portable between workbooks."""
    for row in source:
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            dest = target.cell(cell.row, cell.column, cell.value)
            for attr in ('font', 'fill', 'border', 'alignment', 'protection'):
                setattr(dest, attr, copy(getattr(cell, attr)))
            dest.number_format = cell.number_format
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key] = copy(dim)
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dim)
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    for attr in ('sheet_format', 'sheet_properties', 'page_setup', 'page_margins', 'print_options'):
        setattr(target, attr, copy(getattr(source, attr)))
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    if source.print_area:
        target.print_area = str(source.print_area).split('!')[-1]


def aggregate_details(details):
    groups = {}
    for detail in details:
        code = detail['代码带后缀']
        group = groups.setdefault(code, {'名称': detail['名称'], '代码带后缀': code,
                                        '结算金额差': Decimal(0), 'missing': False})
        if detail.get('结算金额差') is None or detail.get('未匹配完') or detail.get('需人工确认'):
            group['missing'] = True
        else:
            group['结算金额差'] += Decimal(str(detail['结算金额差']))
    return list(groups.values())


def write_reports(result, args, today, find_previous, parse_date, skip_class=False):
    output = Path(args.outdir)
    output.mkdir(parents=True, exist_ok=True)
    prefix = today.strftime('%m%d')
    details = result['逐笔结算金额差']
    support = output / '_support'
    support.mkdir(exist_ok=True)
    (support / f'{prefix}逐笔成本明细.json').write_text(
        json.dumps(details, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

    # Keep the fixed 18-row presentation area; extend only when it is full.
    wb = load_workbook(TEMPLATES / '盈亏情况.xlsx')
    ws = wb.active
    groups = aggregate_details(details)
    bottom = max(18, len(groups) + 1)
    if bottom > 18:
        ws.unmerge_cells('E10:E18')
        for row in range(19, bottom + 1):
            for col in range(1, 6):
                ws.cell(row, col)._style = copy(ws.cell(18, col)._style)
            ws.row_dimensions[row] = copy(ws.row_dimensions[18])
        ws.merge_cells(start_row=10, start_column=5, end_row=bottom, end_column=5)
    for row in range(2, bottom + 1):
        for col in range(1, 5):
            ws.cell(row, col).value = None
    for row, group in enumerate(groups, 2):
        for col, val in enumerate((today, group['名称'], group['代码带后缀'],
                                   None if group['missing'] else float(group['结算金额差'])), 1):
            ws.cell(row, col).value = val
        ws.cell(row, 1).number_format = 'yyyy-mm-dd'
        ws.cell(row, 3).number_format = '@'
        if group['missing']:
            ws.cell(row, 4).value = '待确认成本'
    chai_row = bottom + 2
    # The original funding row can become a bond row when the table grows.
    if chai_row != 20 and len(groups) < 19:
        ws['C20'] = None
        ws['D20'] = None
    ws.cell(chai_row, 3).value = '拆借成本'
    ws.cell(chai_row, 4).value = result['拆借成本（元）']
    ws.cell(chai_row, 4)._style = copy(ws.cell(20, 4)._style)
    incomplete = any(g['missing'] for g in groups)
    ws['E2'] = ('待确认成本' if incomplete else
                f'=SUM(D2:D{max(2, len(groups)+1)})+D{chai_row}')
    ws['E6'] = result['交易量（亿元）']
    ws['E10'] = result['票息收入（万元）']
    first = output / f'{prefix}盈亏情况.xlsx'
    wb.save(first)

    # Roll historical values forward, applying the approved design to the current week only.
    classification = None
    previous_class = getattr(args, 'prev_classification', None) or find_previous('*每日结算分类.xlsx', today)
    if not skip_class:
        master = load_workbook(TEMPLATES / '每日结算分类.xlsx')
        template = master.active
        cls = load_workbook(previous_class, data_only=True) if previous_class else master
        monday = today - datetime.timedelta(days=today.weekday())
        friday = monday + datetime.timedelta(days=4)
        title = f'{monday:%m.%d}-{friday:%m.%d}'
        current = None
        for sh in cls:
            normalized = str(sh['A1'].value or sh.title).replace(' ', '')
            alternatives = [title, f'{monday.month}.{monday.day}-{friday.month}.{friday.day}']
            if any(t in normalized for t in alternatives):
                current = sh
                break
        old_values = [[None] * 5 for _ in range(2)]
        if current is not None:
            old_values = [[current.cell(r, c).value for c in range(2, 7)] for r in (3, 4)]
            index = cls.index(current)
            old_title = current.title
            cls.remove(current)
            current = cls.create_sheet(old_title, index)
        elif cls is master:
            current = cls.active
        else:
            current = cls.create_sheet(title)
        clone_sheet(template, current)
        current.title = title
        current['A1'] = f'{title}   每日结算分类'
        for r, vals in zip((3, 4), old_values):
            for c, val in enumerate(vals, 2):
                current.cell(r, c).value = val
        if today.weekday() > 4:
            raise ValueError('分类模板仅包含周一至周五，请确认周末结算的展示方式')
        current.cell(3, today.weekday()+2).value = '待确认成本' if incomplete else result['代缴（万元）']
        current.cell(4, today.weekday()+2).value = '待确认成本' if incomplete else result['撮合（万元）']
        cls.active = cls.index(current)
        classification = output / f'{prefix}每日结算分类.xlsx'
        cls.save(classification)

    prior_values, prior_date = {}, None
    previous_profit = getattr(args, 'prev_profit', None) or find_previous('*销售交易部当日盈利情况.xlsx', today)
    if previous_profit:
        prev = load_workbook(previous_profit, data_only=True)['最新数据']
        prior_date = parse_date(prev['A1'].value)
        if prior_date is None or prior_date >= today:
            raise ValueError('上一交易日盈利表的最新数据日期必须早于本次计算日')
        prior_values = {prev.cell(r, 1).value: prev.cell(r, 2).value for r in range(3, 14)}
    profit = load_workbook(TEMPLATES / '销售交易部当日盈利情况.xlsx')
    latest, previous = profit['最新数据'], profit['前一日数据']
    for sh, day in ((latest, today), (previous, prior_date)):
        sh['A1'] = day
        year = (day or today).year
        sh['B1'] = datetime.date(year, 12, 31)
        sh['C1'] = datetime.date(year, 1, 1)
    for row in range(3, 14):
        previous.cell(row, 2).value = prior_values.get(previous.cell(row, 1).value)
    value_rows = {3: '当日收益（万元）', 7: '交易量（亿元）', 9: '日均占资（亿元）',
                  10: '时点占资（亿元）', 11: '当日日终持仓（亿元）'}
    for row, key in value_rows.items():
        latest.cell(row, 2).value = None if row == 3 and incomplete else result[key]
    for row, flag in ((5, 'tax_month'), (6, 'tax_year')):
        tax = getattr(args, flag, None)
        latest.cell(row, 2).value = tax if tax is not None else prior_values.get(latest.cell(row, 1).value)
    latest['B4'] = '=IF(AND(ISNUMBER(B3),ISNUMBER(前一日数据!B4)),B3+前一日数据!B4,"待补收益")'
    latest['B8'] = '=IF(ISNUMBER(前一日数据!B8),B7+前一日数据!B8,"待补前日累计")'
    latest['B12'] = '=IF(AND(ISNUMBER(B4),ISNUMBER(B9)),IF(B9=0,"缺数据",B4/B9/10000),"缺数据")'
    latest['B13'] = '=IF(ISNUMBER(B12),B12*365/(A1-C1),"缺数据")'
    profit.active = 0
    third = output / f'{prefix}销售交易部当日盈利情况.xlsx'
    profit.save(third)
    for file in (first, classification, third):
        if file:
            print(f'[生成] {file}')
    if args.tax_month is None or args.tax_year is None:
        print('[注意] 未提供的税费沿用前日；请核实本期贷方与累计贷方，元除以10000写入万元栏。')

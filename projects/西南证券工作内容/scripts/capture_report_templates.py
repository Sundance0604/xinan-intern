"""Capture the user's existing Excel designs; strip all daily business values."""
import argparse
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.packaging.core import DocumentProperties


def capture(source, output):
    output.mkdir(parents=True, exist_ok=True)
    for suffix in ('盈亏情况', '每日结算分类', '销售交易部当日盈利情况'):
        files = list(source.glob(f'*{suffix}.xlsx'))
        if len(files) != 1:
            raise ValueError(f'{suffix}需要且只能有一个参考文件，找到{len(files)}个')
        wb = load_workbook(files[0])
        if suffix == '盈亏情况':
            sh = wb.active
            for row in sh.iter_rows(min_row=2, max_row=18, max_col=4):
                for cell in row:
                    cell.value = None
            sh['E2'] = '=SUM(D2:D18)+D20'
            for address in ('D20', 'E6', 'E10'):
                sh[address] = None
        elif suffix == '每日结算分类':
            sh = wb.active
            # Last existing weekly page is the latest user-approved design.
            candidates = [s for s in wb if s['A3'].value == '代缴' and s['A4'].value == '撮合']
            sh = candidates[-1]
            for other in list(wb):
                if other is not sh:
                    wb.remove(other)
            sh.title = '每周分类'
            sh['A1'] = '每日结算分类'
            for row in sh.iter_rows(min_row=3, max_row=4, min_col=2, max_col=6):
                for cell in row:
                    cell.value = None
            for row in sh.iter_rows(min_row=5):
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.value = None
        else:
            for sh in wb:
                for address in ('A1', 'B1', 'C1', 'I1'):
                    sh[address] = None
                for row in range(3, 14):
                    sh.cell(row, 2).value = None
                if sh['A14'].value:
                    # The supplied long scope note was clipped in its 29pt row.
                    sh.row_dimensions[14].height = max(sh.row_dimensions[14].height or 0, 110)
            sh = wb['最新数据']
            sh['B4'] = '=IF(AND(ISNUMBER(B3),ISNUMBER(前一日数据!B4)),B3+前一日数据!B4,"待补收益")'
            sh['B8'] = '=IF(ISNUMBER(前一日数据!B8),B7+前一日数据!B8,"待补前日累计")'
            sh['B12'] = '=IF(AND(ISNUMBER(B4),ISNUMBER(B9)),IF(B9=0,"缺数据",B4/B9/10000),"缺数据")'
            sh['B13'] = '=IF(ISNUMBER(B12),B12*365/(A1-C1),"缺数据")'
        for sh in wb:
            for row in sh:
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.comment = None
                        cell.hyperlink = None
        wb.properties = DocumentProperties(creator='Sales Reporting', title=f'{suffix}空白模板')
        wb.save(output / f'{suffix}.xlsx')


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('source', type=Path, help='用户确认的三张表所在目录')
    ap.add_argument('--out', type=Path, default=Path(__file__).resolve().parents[1] / 'templates')
    options = ap.parse_args()
    capture(options.source, options.out)

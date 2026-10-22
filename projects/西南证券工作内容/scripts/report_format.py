"""Report presentation only: never round or replace stored numeric values."""
from openpyxl.styles import Font, PatternFill, Alignment

NUMBER = '#,##0.00;(#,##0.00);0.00'
PERCENT = '0.00%'


def format_report(ws, header_row=None):
    ws.sheet_view.showGridLines = False
    for row in ws:
        for cell in row:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = NUMBER
            cell.font = Font(name='Microsoft YaHei', size=11, bold=cell.font.bold)
    if header_row:
        for cell in ws[header_row]:
            cell.fill = PatternFill('solid', fgColor='203864')
            cell.font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[header_row].height = 32
        ws.freeze_panes = f'A{header_row+1}'

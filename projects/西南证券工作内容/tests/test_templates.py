import datetime
import json
import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from template_reports import TEMPLATES, write_reports
from pnl_tables import parse_date


class TemplateTests(unittest.TestCase):
    def result(self, count=2):
        return {'逐笔结算金额差': [
            {'名称': '测试债券', '代码带后缀': '001.SH', '结算金额差': .01},
            {'名称': '测试债券', '代码带后缀': '001.SH', '结算金额差': .02},
        ] if count == 2 else [
            {'名称': f'测试{i}', '代码带后缀': f'{i:06d}.SH', '结算金额差': 1}
            for i in range(count)],
            '拆借成本（元）': -1, '交易量（亿元）': 2.123456,
            '票息收入（万元）': 3.123456, '代缴（万元）': 0,
            '撮合（万元）': .000003, '当日收益（万元）': -.000097,
            '日均占资（亿元）': 1, '时点占资（亿元）': 2, '当日日终持仓（亿元）': 3}

    def generate(self, folder, result, day=datetime.date(2027, 1, 4), **kwargs):
        args = SimpleNamespace(outdir=str(folder), prev_profit=None, prev_classification=None,
                               tax_month=0, tax_year=0)
        for k, v in kwargs.items():
            setattr(args, k, v)
        write_reports(result, args, day, lambda *a: None, parse_date)

    def test_empty_templates_contain_no_business_values(self):
        for file in TEMPLATES.glob('*.xlsx'):
            w = load_workbook(file)
            for s in w:
                for row in s:
                    for cell in row:
                        self.assertFalse(isinstance(cell.value, (int, float, datetime.date)),
                                         (file.name, s.title, cell.coordinate))
        self.assertEqual(len(list(TEMPLATES.glob('*.xlsx'))), 3)

    def test_grouped_output_style_dates_and_missing_prior_guards(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d)
            self.generate(out, self.result())
            s = load_workbook(out/'0104盈亏情况.xlsx').active
            original = load_workbook(TEMPLATES/'盈亏情况.xlsx').active
            self.assertEqual(s['D2'].value, .03)
            self.assertIsNone(s['C3'].value)
            self.assertEqual(s['E2'].value, '=SUM(D2:D2)+D20')
            self.assertEqual(str(s.merged_cells), str(original.merged_cells))
            self.assertEqual(copy(s['A1'].font), copy(original['A1'].font))
            self.assertEqual(s.column_dimensions['B'].width, original.column_dimensions['B'].width)
            self.assertEqual(s['E6'].value, 2.123456)
            c = load_workbook(out/'0104每日结算分类.xlsx').active
            self.assertEqual(c.title, '01.04-01.08')
            self.assertEqual(c['B4'].value, .000003)
            p = load_workbook(out/'0104销售交易部当日盈利情况.xlsx')
            self.assertEqual(p['最新数据']['C1'].value.date(), datetime.date(2027, 1, 1))
            self.assertEqual(p['最新数据']['B1'].value.date(), datetime.date(2027, 12, 31))
            self.assertEqual(p['最新数据']['B5'].value, 0)
            self.assertIn('ISNUMBER(前一日数据!B4)', p['最新数据']['B4'].value)
            self.assertIn('(A1-C1)', p['最新数据']['B13'].value)
            self.assertIsNone(p['前一日数据']['B4'].value)

    def test_overflow_and_new_week_keep_previous_week(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d)
            self.generate(out, self.result(), day=datetime.date(2027, 1, 4))
            before = out/'0104每日结算分类.xlsx'
            self.generate(out, self.result(22), day=datetime.date(2027, 1, 11),
                          prev_classification=str(before))
            s = load_workbook(out/'0111盈亏情况.xlsx').active
            self.assertEqual(s['C23'].value, '000021.SH')
            self.assertEqual(s['C25'].value, '拆借成本')
            self.assertEqual(s['D25'].value, -1)
            self.assertEqual(s['E2'].value, '=SUM(D2:D23)+D25')
            self.assertIn('E10:E23', str(s.merged_cells))
            c = load_workbook(out/'0111每日结算分类.xlsx')
            self.assertEqual(c.sheetnames, ['01.04-01.08', '01.11-01.15'])
            self.assertEqual(c.worksheets[0]['B4'].value, .000003)


if __name__ == '__main__':
    unittest.main()
